import hashlib, json, re
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session
from ..models import Component, RowType, RowTypeComponent

INSTALLATION_RULES = [
    ("5.3", "Podstawy przekładni i łożysk", ["installation of reducer base", "montaż podstawy przekładni"]),
    ("5.4", "Przekładnia mechaniczna", ["installation of slewing reducer", "montaż przekładni"]),
    ("5.5", "Obudowy łożysk", ["installation of bearing housing", "montaż obudowy łożyska"]),
    ("5.6.1", "Rury przy przekładni", ["torque tubes on both sides", "rury po obu stronach przekładni"]),
    ("5.6.2", "Połączenia kolejnych rur", ["diameter reduced torque tubes", "połączenia rur"]),
    ("5.7", "Płatwie", ["installation of purlins", "montaż płatwi"]),
    ("5.8", "Amortyzatory", ["installation of the dampers", "montaż amortyzatorów"]),
    ("5.13", "Zaślepki rur", ["installation of torque tube cover", "montaż zaślepek"]),
]
EXCLUDED = ("wiring", "grounding", "pv module", "pile installation", "ramming", "elektry", "uziem", "wbijan")

def normalize(value):
    return re.sub(r"[^a-z0-9ąćęłńóśźż]+", "_", str(value or "").lower()).strip("_")

def parse_number(value, default=0):
    if isinstance(value, (int, float)): return float(value)
    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", str(value or ""))
    return float(match.group().replace(",", ".")) if match else default

def pdf_analysis(path: Path) -> dict:
    reader = PdfReader(path)
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages)
    lower = text.lower()
    instructions=[]
    for section, title, needles in INSTALLATION_RULES:
        position = next((lower.find(n) for n in needles if lower.find(n) >= 0), -1)
        if position < 0: continue
        excerpt = re.sub(r"\s+", " ", text[position:position + 1200]).strip()
        if any(term in excerpt.lower()[:160] for term in EXCLUDED): continue
        numbered = re.findall(r"(?:^|\s)(\d{1,2})[.)]\s+(.{20,220}?)(?=(?:\s\d{1,2}[.)]\s)|$)", excerpt)
        points = [re.sub(r"\s+", " ", value).strip() for _, value in numbered[:6]]
        instructions.append({"section": section, "title": title, "excerpt": excerpt[:650], "source_points": points})
    torques=[]
    torque_pattern = re.compile(r"(M\d{1,2}(?:\s*[*x×]\s*\d+[^\n,;)]*)?).{0,100}?(?:torque|moment).{0,20}?(\d+(?:[.,]\d+)?)\s*(?:[-~–]\s*(\d+(?:[.,]\d+)?))?\s*N[.·]?m", re.I | re.S)
    for match in torque_pattern.finditer(text):
        item = re.sub(r"\s+", " ", match.group(1)).strip()
        value_min = parse_number(match.group(2)); value_max = parse_number(match.group(3), value_min)
        row={"item":item,"min_nm":value_min,"max_nm":value_max}
        if row not in torques: torques.append(row)
        if len(torques) >= 40: break
    variants=sorted(set(re.findall(r"1\s*[PpXx×*]\s*\d{2,3}", text)))
    return {"text":text,"summary":{"pages":len(pages),"instructions":len(instructions),"torques":len(torques),"variants":variants},"instructions":instructions,"torques":torques}

HEADER_ALIASES = {
    "row_type": {"typ_rzedu","typ_stolu","row_type","tracker_type","wariant","type"},
    "structure": {"typ_konstrukcji","structure_kind","konstrukcja"},
    "length": {"dlugosc","dlugosc_m","length","length_m"},
    "sku": {"sku","indeks","kod","part_no","item_no","nr_katalogowy"},
    "name": {"komponent","nazwa","component","description","opis"},
    "category": {"kategoria","category","grupa"},
    "quantity": {"ilosc","qty","quantity","szt_na_rzad","per_row"},
    "unit": {"jednostka","unit","jm"},
    "torque_min": {"moment_min","torque_min","min_nm"},
    "torque_max": {"moment_max","torque_max","max_nm","moment_nm"},
}

def map_headers(values):
    mapped={}
    for index,value in enumerate(values):
        normalized=normalize(value)
        for key,aliases in HEADER_ALIASES.items():
            if normalized in aliases: mapped[key]=index
    return mapped

def excel_bom_analysis(path: Path, db: Session) -> dict:
    workbook=load_workbook(path, read_only=True, data_only=True)
    imported=0; types=set(); components=set(); warnings=[]
    for sheet in workbook.worksheets:
        rows=list(sheet.iter_rows(values_only=True))
        header_index=-1; headers={}
        for i,row in enumerate(rows[:20]):
            candidate=map_headers(row)
            if "name" in candidate and "quantity" in candidate:
                header_index=i; headers=candidate; break
        if header_index < 0:
            warnings.append(f"{sheet.title}: nie rozpoznano kolumn nazwy i ilości")
            continue
        for row in rows[header_index+1:]:
            def value(key, default=""):
                idx=headers.get(key); return row[idx] if idx is not None and idx < len(row) and row[idx] is not None else default
            name=str(value("name")).strip(); quantity=parse_number(value("quantity"))
            if not name or quantity <= 0: continue
            type_name=str(value("row_type","Typ domyślny")).strip() or "Typ domyślny"
            row_type=db.scalar(select(RowType).where(RowType.name==type_name))
            if not row_type:
                row_type=RowType(name=type_name,structure_kind=str(value("structure","tracker")),reference_length_m=parse_number(value("length")))
                db.add(row_type); db.flush()
            raw_sku=str(value("sku")).strip()
            sku=raw_sku or "AUTO-"+hashlib.sha1(name.lower().encode()).hexdigest()[:10].upper()
            component=db.scalar(select(Component).where(Component.sku==sku))
            if not component:
                tmin=parse_number(value("torque_min"),None); tmax=parse_number(value("torque_max"),tmin)
                component=Component(sku=sku,name=name,category=str(value("category","Inne")),unit=str(value("unit","szt.")),torque_min_nm=tmin,torque_max_nm=tmax)
                db.add(component); db.flush()
            relation=db.scalar(select(RowTypeComponent).where(RowTypeComponent.row_type_id==row_type.id,RowTypeComponent.component_id==component.id))
            if relation: relation.quantity=quantity
            else: db.add(RowTypeComponent(row_type_id=row_type.id,component_id=component.id,quantity=quantity))
            imported+=1; types.add(type_name); components.add(sku)
    db.commit()
    return {"summary":{"sheets":len(workbook.worksheets),"relations":imported,"row_types":len(types),"components":len(components)},"warnings":warnings}

def analyze(path: Path, extension: str, db: Session) -> dict:
    extension=extension.lower()
    if extension==".pdf": return pdf_analysis(path)
    if extension in (".xlsx",".xlsm"): return excel_bom_analysis(path,db)
    if extension==".dwg": return {"summary":{"format":"DWG","status":"zapisany"},"warnings":["DWG zapisano. Do analizy geometrii wyeksportuj rysunek także jako PDF lub DXF."]}
    if extension==".dxf": return {"summary":{"format":"DXF","status":"zapisany"},"warnings":["DXF zapisano; relacje BOM należy dostarczyć w XLSX."]}
    return {"summary":{"status":"zapisany"}}


# --- Analiza plików wgrywanych przez użytkowników -------------------------------------------
# Funkcje poniżej tylko odczytują plik. Nie zapisują niczego do bazy technicznej, więc
# dodatkowy dokument nigdy nie nadpisuje zatwierdzonego BOM-u ani instrukcji.

STATUS_SAVED = "zapisany"
STATUS_ANALYZED = "przeanalizowany"


def _result(status, summary, warnings=None, preview=None):
    return {
        "approved_source": False,
        "approved_for_calculations": False,
        "processing_status": status,
        "summary": summary,
        "warnings": warnings or [],
        "preview": preview or {},
    }


def inspect_pdf_bytes(content: bytes) -> dict:
    from io import BytesIO
    try:
        reader = PdfReader(BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as error:  # uszkodzony lub zaszyfrowany PDF
        return _result(STATUS_SAVED, {"format": "PDF"}, [f"Nie udało się odczytać PDF: {error.__class__.__name__}. Plik zapisano bez analizy."])
    text = "\n".join(pages)
    variants = sorted(set(re.sub(r"\s+", "", value).upper().replace("P", "X").replace("×", "X").replace("*", "X") for value in re.findall(r"1\s*[PpXx×*]\s*\d{2,3}", text)))
    torque_mentions = len(re.findall(r"\d+(?:[.,]\d+)?\s*N[.·]?m", text))
    warnings = ["Dane z pliku pomocniczego nie są używane do obliczeń. Wymagają weryfikacji i zatwierdzenia."]
    if not text.strip():
        warnings.append("PDF nie zawiera warstwy tekstowej (np. skan). Odczytano tylko liczbę stron.")
    return _result(STATUS_ANALYZED, {"format": "PDF", "pages": len(pages), "text_characters": len(text), "variant_labels": variants, "torque_mentions": torque_mentions}, warnings, {"text_excerpt": re.sub(r"\s+", " ", text)[:600]})


def inspect_excel_bytes(content: bytes) -> dict:
    from io import BytesIO
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        return _result(STATUS_SAVED, {"format": "XLSX"}, [f"Nie udało się odczytać arkusza: {error.__class__.__name__}. Plik zapisano bez analizy."])
    sheets = []
    warnings = ["Pozycje z arkusza są tylko podglądem. Nie zmieniają zatwierdzonej bazy BOM."]
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header_index, headers = -1, {}
        for index, row in enumerate(rows[:20]):
            candidate = map_headers(row)
            if "name" in candidate and "quantity" in candidate:
                header_index, headers = index, candidate
                break
        if header_index < 0:
            warnings.append(f"{sheet.title}: nie rozpoznano kolumn nazwy i ilości")
            sheets.append({"name": sheet.title, "rows": len(rows), "bom_rows": 0, "sample": []})
            continue
        sample, count = [], 0
        for row in rows[header_index + 1:]:
            def cell(key):
                position = headers.get(key)
                return row[position] if position is not None and position < len(row) else None
            name = str(cell("name") or "").strip()
            quantity = parse_number(cell("quantity"), None)
            if not name or quantity is None:
                continue
            count += 1
            if len(sample) < 10:
                sample.append({"sku": str(cell("sku") or "").strip(), "name": name, "quantity": quantity, "unit": str(cell("unit") or "").strip()})
        sheets.append({"name": sheet.title, "rows": len(rows), "bom_rows": count, "columns": sorted(headers), "sample": sample})
    workbook.close()
    return _result(STATUS_ANALYZED, {"format": "XLSX", "sheets": len(sheets), "bom_rows": sum(item["bom_rows"] for item in sheets)}, warnings, {"sheets": sheets})


def inspect_dxf_bytes(content: bytes) -> dict:
    if content.startswith(b"AutoCAD Binary DXF"):
        return _result(STATUS_SAVED, {"format": "DXF (binarny)"}, ["Binarny DXF nie jest obsługiwany. Zapisz rysunek jako DXF ASCII lub PDF."])
    text = content.decode("utf-8", errors="replace")
    lines = [line.strip() for line in text.splitlines()]
    if len(lines) < 4 or len(lines) % 2:
        lines = lines[: len(lines) - (len(lines) % 2)]
    pairs = list(zip(lines[0::2], lines[1::2]))
    if not pairs or pairs[0] != ("0", "SECTION"):
        return _result(STATUS_SAVED, {"format": "DXF"}, ["Plik nie wygląda na prawidłowy DXF ASCII. Zapisano bez analizy."])
    section, entity, entities, layers, texts = None, None, {}, set(), []
    expecting_section_name = False
    for code, value in pairs:
        if code == "0":
            if value == "SECTION":
                expecting_section_name = True
                continue
            if value == "ENDSEC":
                section, entity = None, None
                continue
            entity = value
            if section == "ENTITIES":
                entities[value] = entities.get(value, 0) + 1
            continue
        if expecting_section_name and code == "2":
            section, expecting_section_name = value, False
            continue
        if code == "8" and value:
            layers.add(value)
        if section == "ENTITIES" and entity in ("TEXT", "MTEXT") and code == "1" and len(texts) < 50:
            texts.append(value[:200])
    warnings = ["Geometria DXF nie jest przeliczana na BOM. Treści tekstowe wymagają weryfikacji."]
    return _result(STATUS_ANALYZED, {"format": "DXF", "entities": sum(entities.values()), "entity_types": dict(sorted(entities.items())), "layers": len(layers)}, warnings, {"layers": sorted(layers)[:100], "texts": texts})


def inspect_upload(content: bytes, extension: str) -> dict:
    extension = extension.lower()
    if extension == ".pdf":
        return inspect_pdf_bytes(content)
    if extension in (".xlsx", ".xlsm"):
        return inspect_excel_bytes(content)
    if extension == ".dxf":
        return inspect_dxf_bytes(content)
    if extension == ".dwg":
        return _result(STATUS_SAVED, {"format": "DWG"}, ["Aplikacja nie ma parsera DWG. Plik zapisano bez analizy — do odczytu wyeksportuj rysunek do PDF lub DXF."])
    return _result(STATUS_SAVED, {"format": extension.lstrip(".").upper()})
